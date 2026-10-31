from __future__ import annotations

import hashlib
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.shared import Pt
from docxtpl import DocxTemplate, RichText
from jinja2 import Environment, TemplateError, meta, nodes
from jinja2.sandbox import SandboxedEnvironment
from pypdf import PdfReader

from resume_tailor.contracts import (
    ArtifactInfo,
    CandidateFile,
    CandidateRecord,
    CheckResult,
    CliIssue,
    CliResult,
    FactBaseSnapshot,
    SemanticReview,
    SkippedCheck,
    TailoredResumePayload,
    canonical_json_bytes,
)
from resume_tailor.inline_markup import plain_text, to_html_inline, to_markdown_inline, to_rich_text
from resume_tailor.storage import (
    ProtocolError,
    ensure_private_directory,
    read_json_value,
    read_model,
    resolve_managed_path,
    write_exclusive,
    write_json_exclusive,
)
from resume_tailor.validation import check_resume_facts

RENDERER_VERSION = "1"
_SAFE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,79}$")
_ALLOWED_FILTERS = {"default", "join", "length", "lower", "replace", "trim", "upper"}
_UNRENDERED_TAG = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.DOTALL)


def _default_template_bytes() -> bytes:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)
    document.add_paragraph("{{ name }}")
    document.add_paragraph("{{ contact }}")
    document.add_paragraph("{{r summary }}")
    document.add_paragraph("SKILLS")
    document.add_paragraph("{%p for skill in skills %}")
    document.add_paragraph("{{ skill }}")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("{%p for experience in experiences %}")
    document.add_paragraph("{{ experience.role }} | {{ experience.company }}")
    document.add_paragraph("{{ experience.dates }}")
    document.add_paragraph("{%p for bullet in experience.bullets %}")
    document.add_paragraph("{{r bullet }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("{%p for project in projects %}")
    document.add_paragraph("{{ project.name }} | {{ project.role }}")
    document.add_paragraph("{%p for bullet in project.bullets %}")
    document.add_paragraph("{{r bullet }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("{%p for education in educations %}")
    document.add_paragraph("{{ education.label }}")
    document.add_paragraph("{%p endfor %}")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _plain_context(value: Any) -> bool:
    if value is None or isinstance(value, (str, int, float, bool, RichText)):
        return True
    if isinstance(value, list):
        return all(_plain_context(item) for item in value)
    if isinstance(value, dict):
        return all(isinstance(key, str) and _plain_context(item) for key, item in value.items())
    return False


def _validate_template(template_bytes: bytes, context: dict[str, Any]) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(template_bytes)) as archive:
            part_names = [
                "word/document.xml",
                *sorted(
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
                ),
            ]
            sources = [
                archive.read(name)
                .decode("utf-8")
                .replace("{%p", "{%")
                .replace("{%tr", "{%")
                .replace("{{r", "{{")
                for name in part_names
            ]
        environment = Environment()
        environment.globals.clear()
        syntax_tree = environment.parse("\n".join(sources))
    except (KeyError, UnicodeError, TemplateError, zipfile.BadZipFile) as error:
        raise ProtocolError(
            "DOCX template failed static validation.", code="template_invalid"
        ) from error

    unknown_variables = meta.find_undeclared_variables(syntax_tree) - set(context)
    unsafe_attributes = any(
        node.attr.startswith("_") for node in syntax_tree.find_all(nodes.Getattr)
    )
    unsafe_items = any(
        isinstance(node.arg, nodes.Const)
        and isinstance(node.arg.value, str)
        and node.arg.value.startswith("_")
        for node in syntax_tree.find_all(nodes.Getitem)
    )
    unsafe_calls = any(True for _ in syntax_tree.find_all(nodes.Call))
    unknown_filters = {
        node.name
        for node in syntax_tree.find_all(nodes.Filter)
        if node.name not in _ALLOWED_FILTERS
    }
    unsafe_tags = any(
        isinstance(
            node,
            (nodes.Extends, nodes.FromImport, nodes.Import, nodes.Include, nodes.Macro),
        )
        for node in syntax_tree.find_all(nodes.Node)
    )
    if (
        unknown_variables
        or unsafe_attributes
        or unsafe_items
        or unsafe_calls
        or unknown_filters
        or unsafe_tags
    ):
        raise ProtocolError("DOCX template failed static validation.", code="template_invalid")


def _sandboxed_template_environment() -> SandboxedEnvironment:
    environment = SandboxedEnvironment(autoescape=True)
    environment.globals.clear()
    environment.filters = {
        name: environment.filters[name] for name in sorted(_ALLOWED_FILTERS)
    }
    return environment


def render_docx_template(template_bytes: bytes, context: dict[str, Any]) -> bytes:
    if not _plain_context(context):
        raise ProtocolError(
            "Template context must contain plain data only.", code="template_invalid"
        )
    _validate_template(template_bytes, context)
    try:
        template = DocxTemplate(io.BytesIO(template_bytes))
        environment = _sandboxed_template_environment()
        template.render(context, jinja_env=environment, autoescape=True)
        stream = io.BytesIO()
        template.save(stream)
        return stream.getvalue()
    except Exception as error:
        raise ProtocolError(
            "DOCX template rendering failed safely.", code="template_invalid"
        ) from error


def _template_bytes(template_name: str | None) -> tuple[bytes, str]:
    name = template_name or "default"
    if name == "default":
        builtin_default = Path(__file__).parent / "templates_builtin" / "default.docx"
        if builtin_default.is_file():
            return builtin_default.read_bytes(), "default"
        return _default_template_bytes(), "default"
    if not _SAFE_ID.fullmatch(name):
        raise ProtocolError("Unknown template name.", code="unknown_template")
    builtin_path = Path(__file__).parent / "templates_builtin" / f"{name}.docx"
    if builtin_path.is_file():
        return builtin_path.read_bytes(), name
    try:
        path = resolve_managed_path(
            "private_data/storage", Path("templates_custom") / f"{name}.docx"
        )
        if not path.is_file():
            raise ProtocolError("Unknown template name.", code="unknown_template")
        return path.read_bytes(), name
    except ProtocolError:
        raise
    except OSError as error:
        raise ProtocolError("Could not read the selected DOCX template.") from error


def _resume_data(
    snapshot: FactBaseSnapshot,
    payload: TailoredResumePayload,
) -> dict[str, Any]:
    fact_base = snapshot.fact_base
    skills = {item.skill_id: item.skill_name for item in fact_base.verified_skills}
    experiences = {item.exp_id: item for item in fact_base.work_experiences}
    projects = {item.project_id: item for item in fact_base.projects}
    educations = {item.edu_id: item for item in fact_base.education}
    experience_sections = []
    for section in payload.experiences:
        fact = experiences.get(section.exp_id)
        if fact is not None:
            experience_sections.append(
                {
                    "company": (
                        f"{fact.company_name} ({fact.location})"
                        if fact.location
                        else fact.company_name
                    ),
                    "role": fact.role_title,
                    "dates": f"{fact.start_date} - {fact.end_date}",
                    "bullets": [bullet.rewritten_text for bullet in section.selected_bullets],
                }
            )
    project_sections = []
    for section in payload.projects:
        fact = projects.get(section.project_id)
        if fact is not None:
            project_sections.append(
                {
                    "name": fact.project_name,
                    "role": fact.role_title,
                    "bullets": [bullet.rewritten_text for bullet in section.selected_bullets],
                }
            )
    education_sections = []
    for education_id in payload.selected_education_ids:
        fact = educations.get(education_id)
        if fact is not None:
            label = " | ".join(
                value
                for value in (fact.degree, fact.field_of_study, fact.institution, fact.end_year)
                if value
            )
            education_sections.append({"label": label})
    identity = fact_base.identity
    contact = " | ".join(
        value
        for value in (
            identity.email,
            identity.phone,
            identity.location,
            identity.linkedin_url,
            identity.github_url,
            identity.portfolio_url,
        )
        if value
    )
    return {
        "name": identity.full_name,
        "contact": contact,
        "summary": payload.summary_statement or "",
        "skills": [skills[value] for value in payload.selected_skill_ids if value in skills],
        "experiences": experience_sections,
        "projects": project_sections,
        "educations": education_sections,
    }


def _docx_context(data: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": data["name"],
        "contact": data["contact"],
        "summary": to_rich_text(data["summary"]),
        "skills": data["skills"],
        "experiences": [
            {
                **experience,
                "bullets": [to_rich_text(value) for value in experience["bullets"]],
            }
            for experience in data["experiences"]
        ],
        "projects": [
            {
                **project,
                "bullets": [to_rich_text(value) for value in project["bullets"]],
            }
            for project in data["projects"]
        ],
        "educations": data["educations"],
    }


def markdown_document(data: dict[str, Any]) -> str:
    lines: list[str] = []
    if data["name"]:
        lines.extend([f"# {to_markdown_inline(data['name'])}", ""])
    if data["contact"]:
        lines.extend([to_markdown_inline(data["contact"]), ""])
    if data["summary"]:
        lines.extend(["## Summary", "", to_markdown_inline(data["summary"]), ""])
    for title, entries in (
        ("Skills", data["skills"]),
        ("Experience", data["experiences"]),
        ("Projects", data["projects"]),
        ("Education", data["educations"]),
    ):
        if not entries:
            continue
        lines.extend([f"## {title}", ""])
        if title == "Skills":
            lines.extend(f"- {to_markdown_inline(value)}" for value in entries)
        elif title in {"Experience", "Projects"}:
            for entry in entries:
                heading = " | ".join(
                    value
                    for value in (entry.get("role"), entry.get("company"), entry.get("name"))
                    if value
                )
                if entry.get("dates"):
                    heading = f"{heading} | {entry['dates']}"
                lines.extend([f"### {to_markdown_inline(heading)}", ""])
                lines.extend(f"- {to_markdown_inline(value)}" for value in entry["bullets"])
                lines.append("")
        else:
            lines.extend(f"- {to_markdown_inline(entry['label'])}" for entry in entries)
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def markdown_html(data: dict[str, Any]) -> str:
    parts = ['<!doctype html><html><head><meta charset="utf-8"></head><body>']
    if data["name"]:
        parts.append(f"<h1>{to_html_inline(data['name'])}</h1>")
    if data["contact"]:
        parts.append(f"<p>{to_html_inline(data['contact'])}</p>")
    if data["summary"]:
        parts.extend(["<h2>Summary</h2>", f"<p>{to_html_inline(data['summary'])}</p>"])
    for title, entries in (
        ("Skills", data["skills"]),
        ("Experience", data["experiences"]),
        ("Projects", data["projects"]),
        ("Education", data["educations"]),
    ):
        if not entries:
            continue
        parts.append(f"<h2>{title}</h2>")
        if title == "Skills":
            items = "".join(f"<li>{to_html_inline(value)}</li>" for value in entries)
            parts.append(f"<ul>{items}</ul>")
        elif title in {"Experience", "Projects"}:
            for entry in entries:
                heading = " | ".join(
                    value
                    for value in (entry.get("role"), entry.get("company"), entry.get("name"))
                    if value
                )
                if entry.get("dates"):
                    heading = f"{heading} | {entry['dates']}"
                parts.append(f"<h3>{to_html_inline(heading)}</h3>")
                items = "".join(
                    f"<li>{to_html_inline(value)}</li>" for value in entry["bullets"]
                )
                parts.append(f"<ul>{items}</ul>")
        else:
            items = "".join(
                f"<li>{to_html_inline(entry['label'])}</li>" for entry in entries
            )
            parts.append(f"<ul>{items}</ul>")
    parts.append("</body></html>")
    result = "".join(parts)
    if re.search(r"<\s*(img|a|link|script|style)\b|\b(src|href)\s*=", result, re.I):
        raise ProtocolError("Generated HTML contains a disallowed resource reference.")
    return result


def _find_soffice() -> str | None:
    executable = shutil.which("soffice")
    if executable is not None:
        return executable
    if sys.platform == "win32" or os.name == "nt":
        for program_files in (
            os.environ.get("ProgramFiles", r"C:\Program Files"),
            os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
        ):
            candidate = Path(program_files) / "LibreOffice" / "program" / "soffice.exe"
            if candidate.is_file():
                return str(candidate)
    elif sys.platform == "darwin":
        mac_app = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if mac_app.is_file():
            return str(mac_app)
    return None


def _convert_to_pdf(source_name: str, source_bytes: bytes) -> bytes | None:
    executable = _find_soffice()
    if executable is None:
        return None
    scratch_root = ensure_private_directory("private_data/storage/tmp/render_scratch")
    try:
        with tempfile.TemporaryDirectory(prefix="resume-render-", dir=scratch_root) as temp_name:
            work_dir = Path(temp_name)
            source_path = write_exclusive(work_dir / source_name, source_bytes)
            output_dir = work_dir / "output"
            profile_dir = work_dir / "profile"
            output_dir.mkdir(mode=0o700)
            profile_dir.mkdir(mode=0o700)
            command = [
                executable,
                "--headless",
                "--nologo",
                "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(source_path),
            ]
            completed = subprocess.run(
                command,
                check=False,
                capture_output=True,
                timeout=120,
            )
            pdf_path = output_dir / f"{source_path.stem}.pdf"
            if completed.returncode != 0 or not pdf_path.is_file():
                raise ProtocolError(
                    "LibreOffice could not produce a PDF.", code="pdf_render_failed"
                )
            return pdf_path.read_bytes()
    except subprocess.TimeoutExpired as error:
        raise ProtocolError(
            "LibreOffice timed out while producing a PDF.", code="pdf_timeout"
        ) from error
    except OSError as error:
        raise ProtocolError(
            "Could not run LibreOffice for PDF conversion.", code="pdf_render_failed"
        ) from error


def _normalize_pdf_text(value: str) -> str:
    return " ".join(value.split()).casefold()


def verify_pdf_health(
    pdf_bytes: bytes,
    data: dict[str, Any],
    target_pages: int | None,
) -> list[CheckResult]:
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes), strict=True)
        pages = list(reader.pages)
        extracted = "\n".join(page.extract_text() or "" for page in pages)
    except Exception:
        return [
            CheckResult(
                check="pdf_readable",
                passed=False,
                issue_code="pdf_unreadable",
                message="Generated PDF could not be inspected.",
            )
        ]

    within_limit = bool(pages) and (target_pages is None or len(pages) <= target_pages)
    has_unrendered_tags = _UNRENDERED_TAG.search(extracted) is not None
    checks = [
        CheckResult(
            check="pdf_page_count",
            passed=within_limit,
            issue_code=None if within_limit else "pdf_page_limit_exceeded",
            message="PDF page count is within the selected limit."
            if within_limit
            else "PDF page count is empty or exceeds the selected limit.",
        ),
        CheckResult(
            check="pdf_text_layer",
            passed=bool(extracted.strip()),
            issue_code=None if extracted.strip() else "pdf_text_layer_missing",
            message="PDF has an extractable text layer."
            if extracted.strip()
            else "PDF has no extractable text layer.",
        ),
        CheckResult(
            check="pdf_unrendered_tags",
            passed=not has_unrendered_tags,
            issue_code="unrendered_template_tag" if has_unrendered_tags else None,
            message="PDF contains no unrendered template tags."
            if not has_unrendered_tags
            else "PDF contains an unrendered template tag.",
        ),
        CheckResult(
            check="pdf_missing_glyphs",
            passed="\ufffd" not in extracted,
            issue_code=None if "\ufffd" not in extracted else "pdf_missing_glyph",
            message="PDF text contains no replacement glyphs."
            if "\ufffd" not in extracted
            else "PDF text contains a replacement glyph.",
        ),
    ]
    geometry_valid = all(
        float(page.mediabox.width) > 0 and float(page.mediabox.height) > 0 for page in pages
    )
    checks.append(
        CheckResult(
            check="pdf_page_geometry",
            passed=geometry_valid,
            issue_code=None if geometry_valid else "pdf_invalid_page_geometry",
            message="PDF page dimensions are valid."
            if geometry_valid
            else "PDF contains invalid page dimensions.",
        )
    )
    normalized_pdf = _normalize_pdf_text(extracted)
    expected_texts = ([data["summary"]] if data["summary"] else []) + [
        bullet for section in data["experiences"] for bullet in section["bullets"]
    ] + [bullet for section in data["projects"] for bullet in section["bullets"]]
    for index, value in enumerate(expected_texts, start=1):
        expected = _normalize_pdf_text(plain_text(value))
        present = not expected or expected in normalized_pdf
        checks.append(
            CheckResult(
                check="selected_text_present",
                passed=present,
                issue_code=None if present else "selected_text_missing",
                target_id=f"rendered_text_{index}",
                message="Selected resume text appears in the PDF."
                if present
                else "Selected resume text is missing from the PDF text layer.",
            )
        )
    return checks


def _normal_label(value: str) -> str:
    return " ".join(value.split()).casefold()


def _resume_filename_stem(
    company_name: str | None,
    role_title: str | None,
    created_at: datetime,
) -> str:
    label = " ".join(value for value in (company_name, role_title) if value)
    normalized = unicodedata.normalize("NFKD", label).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "-", normalized.casefold()).strip("-") or "resume"
    slug = slug[:180].rstrip("-") or "resume"
    timestamp = created_at.astimezone().strftime("%Y%m%d-%H%M%S")
    return f"{slug}-{timestamp}"


def _company_slug(company_name: str) -> str:
    normalized = unicodedata.normalize("NFKD", company_name).encode("ascii", "ignore").decode()
    label = re.sub(r"[^a-z0-9]+", "-", normalized.casefold()).strip("-") or "company"
    suffix = hashlib.sha256(_normal_label(company_name).encode("utf-8")).hexdigest()[:10]
    return f"{label[:42]}-{suffix}"


def _job_metadata(path: Path) -> dict[str, Any]:
    metadata_path = path / "job.json"
    if metadata_path.is_symlink() or not metadata_path.is_file():
        raise ProtocolError("Managed job metadata is missing or unsafe.", code="unsafe_path")
    value = read_json_value(metadata_path)
    if not isinstance(value, dict):
        raise ProtocolError("Managed job metadata is invalid.")
    if value.get("job_id") != path.name:
        raise ProtocolError("Managed job metadata is inconsistent.")
    return value


def _managed_candidate_directory(
    company_name: str,
    role_title: str,
    requested_job_id: str | None,
    run_id: str,
) -> tuple[Path, str, str]:
    if not _SAFE_ID.fullmatch(run_id):
        raise ProtocolError("Run ID contains unsupported characters.", code="invalid_run_id")
    company_dir = ensure_private_directory(
        Path("private_data/storage/companies") / _company_slug(company_name)
    )
    jobs_dir = ensure_private_directory(company_dir / "jobs")
    role_key = _normal_label(role_title)
    company_key = _normal_label(company_name)
    matching: list[tuple[str, Path]] = []
    for job_path in jobs_dir.iterdir():
        if job_path.is_symlink():
            raise ProtocolError("Managed storage contains a symbolic link.", code="unsafe_path")
        if not job_path.is_dir():
            continue
        metadata = _job_metadata(job_path)
        same_company = _normal_label(str(metadata.get("company_name", ""))) == company_key
        same_role = _normal_label(str(metadata.get("role_title", ""))) == role_key
        if same_company and same_role:
            matching.append((str(metadata["job_id"]), job_path))

    if requested_job_id is not None:
        if not _SAFE_ID.fullmatch(requested_job_id):
            raise ProtocolError("Job ID contains unsupported characters.", code="invalid_job_id")
        existing = jobs_dir / requested_job_id
        if existing.exists():
            metadata = _job_metadata(existing)
            if _normal_label(str(metadata.get("company_name", ""))) != company_key or _normal_label(
                str(metadata.get("role_title", ""))
            ) != role_key:
                raise ProtocolError("Job ID does not match the confirmed company and role.")
            job_id = requested_job_id
            job_dir = existing
        else:
            job_id = requested_job_id
            job_dir = ensure_private_directory(existing)
            write_json_exclusive(
                job_dir / "job.json",
                {
                    "schema_version": 1,
                    "job_id": job_id,
                    "company_name": company_name,
                    "role_title": role_title,
                },
            )
    elif len(matching) > 1:
        raise ProtocolError("Multiple matching jobs exist; pass --job-id explicitly.")
    elif matching:
        job_id, job_dir = matching[0]
    else:
        job_id = f"job_{uuid4().hex[:12]}"
        job_dir = ensure_private_directory(jobs_dir / job_id)
        write_json_exclusive(
            job_dir / "job.json",
            {
                "schema_version": 1,
                "job_id": job_id,
                "company_name": company_name,
                "role_title": role_title,
            },
        )

    run_dir = ensure_private_directory(job_dir / "runs" / run_id)
    artifact_id = f"artifact_{uuid4().hex}"
    candidate_dir = run_dir / artifact_id
    try:
        candidate_dir.mkdir(mode=0o700, exist_ok=False)
    except FileExistsError as error:
        raise ProtocolError(
            "Candidate directory already exists.", code="output_conflict"
        ) from error
    except OSError as error:
        raise ProtocolError("Could not create the managed candidate directory.") from error
    return candidate_dir, job_id, artifact_id


def _export_candidate_directory(path: str | Path) -> Path:
    target = Path(path).expanduser()
    if "\0" in str(target) or not target.parent.is_dir():
        raise ProtocolError("Export parent directory must already exist.")
    try:
        target.mkdir(mode=0o700, exist_ok=False)
    except FileExistsError as error:
        raise ProtocolError("Output target already exists.", code="output_conflict") from error
    except OSError as error:
        raise ProtocolError("Could not create the selected export directory.") from error
    return target.resolve(strict=True)


def _confirmation_summary(snapshot: FactBaseSnapshot) -> dict[str, Any]:
    items = [
        snapshot.fact_base.identity,
        *snapshot.fact_base.education,
        *snapshot.fact_base.work_experiences,
        *snapshot.fact_base.projects,
        *snapshot.fact_base.verified_skills,
        *[
            bullet
            for experience in snapshot.fact_base.work_experiences
            for bullet in experience.bullets
        ],
        *[
            bullet
            for project in snapshot.fact_base.projects
            for bullet in project.bullets
        ],
    ]
    return {
        "snapshot_sha256": snapshot.snapshot_sha256,
        "ledger_sha256": snapshot.ledger_sha256,
        "confirmed_item_count": sum(item.user_confirmed for item in items),
        "new_or_modified_facts_in_payload": [],
    }


def _artifact_info(directory: Path, filename: str, role: str, content: bytes) -> ArtifactInfo:
    return ArtifactInfo(
        path=str(directory / filename),
        role=role,
        sha256=hashlib.sha256(content).hexdigest(),
    )


def render_resume(
    *,
    snapshot_path: str | Path,
    payload_path: str | Path,
    company_name: str | None,
    role_title: str | None,
    out_dir: str | Path | None,
    job_id: str | None,
    review_path: str | Path | None,
    template_name: str | None,
    output_format: str | None,
    run_id: str | None,
    requirements_path: str | Path | None = None,
) -> tuple[CliResult, int]:
    report = check_resume_facts(
        snapshot_path,
        payload_path,
        requirements_path=requirements_path,
        review_path=review_path,
    )
    if any(not check.passed for check in report.checks):
        return (
            CliResult(
                command="render",
                status="blocked",
                issues=[
                    CliIssue(
                        code=check.issue_code or check.check,
                        severity="error",
                        message=check.message,
                    )
                    for check in report.checks
                    if not check.passed
                ],
                skipped_checks=report.skipped_checks,
                data={"audit_report": report.model_dump(mode="json")},
            ),
            3,
        )

    snapshot = read_model(snapshot_path, FactBaseSnapshot)
    payload = read_model(payload_path, TailoredResumePayload)
    review = read_model(review_path, SemanticReview) if review_path is not None else None
    review_current = review is not None and review.payload_sha256 == report.payload_sha256
    review_mode = review.mode if review_current and review is not None else "NONE"
    review_status = review.status if review_current and review is not None else "NOT_RUN"
    actual_run_id = run_id or f"run_{uuid4().hex}"
    if not _SAFE_ID.fullmatch(actual_run_id):
        raise ProtocolError("Run ID contains unsupported characters.", code="invalid_run_id")
    if out_dir is not None and Path(out_dir).expanduser().exists():
        raise ProtocolError("Output target already exists.", code="output_conflict")

    created_at = datetime.now(UTC)
    filename_stem = _resume_filename_stem(company_name, role_title, created_at)
    docx_filename = f"{filename_stem}.docx"
    data = _resume_data(snapshot, payload)
    effective_format = (output_format or payload.run_config.render_format).casefold()
    if effective_format not in {"docx", "markdown"}:
        raise ProtocolError("Unsupported resume output format.")
    rendered_files: list[tuple[str, str, bytes]] = []
    issues: list[CliIssue] = []
    template_fallback = False
    selected_template = template_name or payload.run_config.template_id or "default"
    if effective_format == "docx":
        template_content, template_id = _template_bytes(selected_template)
        try:
            docx_content = render_docx_template(template_content, _docx_context(data))
            rendered_files.append((docx_filename, "resume_docx", docx_content))
        except ProtocolError as error:
            if error.code != "template_invalid" or template_id == "default":
                raise
            template_fallback = True
            effective_format = "markdown"
            report.skipped_checks.append(
                SkippedCheck(
                    check="docx_template",
                    reason="Selected DOCX template failed validation.",
                )
            )
            issues.append(
                CliIssue(
                    code="template_fallback",
                    severity="warning",
                    message="The DOCX template was rejected; rendered the safe Markdown fallback.",
                )
            )
    elif template_name not in (None, "default"):
        _template_bytes(selected_template)

    if effective_format == "markdown":
        markdown = markdown_document(data).encode("utf-8")
        html_content = markdown_html(data).encode("utf-8")
        rendered_files.append((f"{filename_stem}.md", "resume_markdown", markdown))
        pdf_source_name = f"{filename_stem}.html"
        pdf_source = html_content
        render_path_base = "markdown"
    else:
        pdf_source_name = docx_filename
        pdf_source = rendered_files[0][2]
        render_path_base = "docx"

    pdf_content = _convert_to_pdf(pdf_source_name, pdf_source)
    pdf_health_passed = False
    if pdf_content is not None:
        rendered_files.append((f"{filename_stem}.pdf", "resume_pdf", pdf_content))
        health_checks = verify_pdf_health(pdf_content, data, payload.run_config.target_pages)
        report.checks.extend(health_checks)
        pdf_health_passed = all(check.passed for check in health_checks)
    else:
        report.skipped_checks.append(
            SkippedCheck(
                check="pdf_generation_and_health",
                reason=(
                    "LibreOffice (soffice) is unavailable. No PDF was generated or health-checked; "
                    "the DOCX or Markdown candidate is retained for manual conversion."
                ),
            )
        )
        issues.append(
            CliIssue(
                code="libreoffice_unavailable",
                severity="warning",
                message="No PDF was generated; the candidate is not a complete PDF export.",
            )
        )
    report.skipped_checks.extend(
        [
            SkippedCheck(
                check="character_budget",
                reason="No measured budget is available for this template and output language.",
            ),
            SkippedCheck(
                check="visual_layout_review",
                reason=(
                    "Automated checks cannot establish absence of clipping or overlap; "
                    "preview is required."
                ),
            ),
        ]
    )
    render_path = (
        f"{render_path_base}_libreoffice_pdf"
        if pdf_content
        else f"{render_path_base}_only_no_pdf"
    )
    if template_fallback:
        render_path = (
            "markdown_template_fallback_libreoffice_pdf"
            if pdf_content
            else "markdown_template_fallback_only_no_pdf"
        )
    report.render_path = render_path
    deterministic_passed = all(check.passed for check in report.checks)
    report.approval_conditions.update(
        {
            "deterministic_checks_passed": deterministic_passed,
            "render_health_passed": pdf_health_passed,
            "semantic_review_mode": review_mode,
            "semantic_review_status": review_status,
            "semantic_review_current": bool(review_current),
            "semantic_review_waived": False,
            "requires_explicit_user_waiver": review_status in {"NOT_RUN", "UNAVAILABLE"},
            "can_approve": deterministic_passed and pdf_health_passed and review_status == "PASSED",
            "can_publish": deterministic_passed and pdf_health_passed and review_status == "PASSED",
            "confirmation_summary": _confirmation_summary(snapshot),
        }
    )

    failed_checks = [check for check in report.checks if not check.passed]
    if failed_checks:
        issues.extend(
            CliIssue(
                code=check.issue_code or check.check,
                severity="error",
                message=check.message,
            )
            for check in failed_checks
        )
        return (
            CliResult(
                command="render",
                status="blocked",
                issues=issues,
                skipped_checks=report.skipped_checks,
                data={
                    "render_path": render_path,
                    "pdf_generated": pdf_content is not None,
                    "confirmation_summary": _confirmation_summary(snapshot),
                    "audit_report": report.model_dump(mode="json"),
                },
            ),
            3,
        )

    if review_current and review is not None:
        rendered_files.append(
            ("semantic_review.json", "semantic_review", canonical_json_bytes(review) + b"\n")
        )
    rendered_files.append(("payload.json", "payload", canonical_json_bytes(payload) + b"\n"))
    rendered_files.append(
        ("audit_report.json", "audit_report", canonical_json_bytes(report) + b"\n")
    )

    if company_name is not None and role_title is not None:
        candidate_dir, actual_job_id, artifact_id = _managed_candidate_directory(
            company_name,
            role_title,
            job_id,
            actual_run_id,
        )
    elif out_dir is not None:
        candidate_dir = _export_candidate_directory(out_dir)
        actual_job_id = job_id or f"job_{uuid4().hex[:12]}"
        artifact_id = f"artifact_{uuid4().hex}"
    else:
        raise ProtocolError("A managed job or explicit export directory is required.")

    candidate_files = [
        CandidateFile(
            path=filename,
            role=role,
            sha256=hashlib.sha256(content).hexdigest(),
        )
        for filename, role, content in rendered_files
    ]
    candidate = CandidateRecord(
        job_id=actual_job_id,
        run_id=actual_run_id,
        artifact_id=artifact_id,
        company_name=company_name,
        role_title=role_title,
        payload_sha256=report.payload_sha256,
        snapshot_sha256=snapshot.snapshot_sha256,
        ledger_sha256=snapshot.ledger_sha256,
        job_requirements_sha256=payload.job_requirements_sha256,
        files=candidate_files,
        review_mode=review_mode,
        review_status=review_status,
        render_path=render_path,
        created_at=created_at,
    )

    created_paths: list[Path] = []
    try:
        for filename, _, content in rendered_files:
            created_paths.append(write_exclusive(candidate_dir / filename, content))
        candidate_path = write_json_exclusive(candidate_dir / "candidate.json", candidate)
        created_paths.append(candidate_path)
    except ProtocolError:
        for path in created_paths:
            path.unlink(missing_ok=True)
        try:
            candidate_dir.rmdir()
        except OSError:
            pass
        raise

    artifacts = [
        _artifact_info(candidate_dir, filename, role, content)
        for filename, role, content in rendered_files
    ]
    artifacts.append(
        _artifact_info(
            candidate_dir,
            "candidate.json",
            "candidate_record",
            candidate_path.read_bytes(),
        )
    )
    final_checks_passed = all(check.passed for check in report.checks)
    if not final_checks_passed:
        status = "blocked"
        exit_code = 3
        issues.extend(
            CliIssue(
                code=check.issue_code or check.check,
                severity="error",
                message=check.message,
            )
            for check in report.checks
            if not check.passed
        )
    elif pdf_content is None:
        status = "degraded"
        exit_code = 4
    else:
        status = "ok"
        exit_code = 0
    return (
        CliResult(
            command="render",
            status=status,
            issues=issues,
            artifacts=artifacts,
            skipped_checks=report.skipped_checks,
            data={
                "job_id": actual_job_id,
                "run_id": actual_run_id,
                "artifact_id": artifact_id,
                "candidate_dir": str(candidate_dir),
                "render_path": render_path,
                "pdf_generated": pdf_content is not None,
                "confirmation_summary": _confirmation_summary(snapshot),
                "audit_report": report.model_dump(mode="json"),
            },
        ),
        exit_code,
    )