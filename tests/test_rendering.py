import io
from pathlib import Path

import pytest
from docx import Document

from resume_tailor.contracts import (
    CheckResult,
    EvidenceBullet,
    FactBaseSnapshot,
    Identity,
    JobKeyword,
    JobRequirements,
    RewrittenBullet,
    TailoredExperienceSection,
    TailoredResumePayload,
    UserFactBase,
    WorkExperienceFact,
    content_sha256,
    sha256_json,
)
from resume_tailor.inline_markup import plain_text, to_html_inline, to_markdown_inline, to_rich_text
from resume_tailor.rendering import (
    _default_template_bytes,
    markdown_html,
    render_docx_template,
    render_resume,
)
from resume_tailor.storage import ProtocolError


def test_inline_markup_escapes_injection_and_external_resources() -> None:
    value = '![x](http://example.test) <script>alert(1)</script> [x](javascript:alert(1))'

    rendered = to_html_inline(value)
    markdown = to_markdown_inline(value)

    assert "<script>" not in rendered
    assert "<img" not in rendered
    assert "<a" not in rendered
    assert "&lt;script&gt;" in rendered
    assert r"\!\[x\]\(http://example\.test\)" in markdown
    assert "&lt;script&gt;" in markdown


def test_only_supported_bold_and_italic_markers_are_interpreted() -> None:
    value = "**bold** and *italic* with C*"

    assert to_html_inline(value) == "<strong>bold</strong> and <em>italic</em> with C*"
    assert to_markdown_inline(value) == r"**bold** and *italic* with C\*"
    assert plain_text(value) == "bold and italic with C*"


def test_default_docx_template_renders_inside_sandbox() -> None:
    context = {
        "name": "Synthetic Engineer",
        "contact": "",
        "summary": to_rich_text(""),
        "skills": [],
        "experiences": [
            {
                "company": "Synthetic Systems",
                "role": "Engineer",
                "dates": "2020 - 2022",
                "bullets": [to_rich_text("Built a reliable service.")],
            }
        ],
        "projects": [],
        "educations": [],
    }

    rendered = render_docx_template(_default_template_bytes(), context)
    document = Document(io.BytesIO(rendered))

    assert "Synthetic Systems" in " ".join(paragraph.text for paragraph in document.paragraphs)
    document_text = " ".join(paragraph.text for paragraph in document.paragraphs)
    assert "Built a reliable service." in document_text


def test_ssti_probe_template_is_rejected_without_producing_a_document() -> None:
    document = Document()
    document.add_paragraph("{{ cycler.__init__.__globals__ }}")
    template = io.BytesIO()
    document.save(template)

    with pytest.raises(ProtocolError, match="failed static validation"):
        render_docx_template(template.getvalue(), {"name": "Synthetic"})


@pytest.mark.parametrize("part_name", ["header", "footer"])
def test_docx_header_and_footer_templates_are_statically_validated(part_name: str) -> None:
    document = Document()
    part = getattr(document.sections[0], part_name)
    part.paragraphs[0].text = "{{ name | safe }}"
    template = io.BytesIO()
    document.save(template)

    with pytest.raises(ProtocolError, match="failed static validation"):
        render_docx_template(template.getvalue(), {"name": "Synthetic"})


def test_docx_templates_render_only_allowlisted_jinja_filters() -> None:
    from resume_tailor.rendering import _sandboxed_template_environment

    document = Document()
    document.add_paragraph("{{ name | upper }}")
    template = io.BytesIO()
    document.save(template)
    environment = _sandboxed_template_environment()

    assert set(environment.filters) == {
        "default",
        "join",
        "length",
        "lower",
        "replace",
        "trim",
        "upper",
    }
    assert not environment.globals
    rendered = render_docx_template(template.getvalue(), {"name": "Synthetic"})
    assert "SYNTHETIC" in " ".join(
        paragraph.text for paragraph in Document(io.BytesIO(rendered)).paragraphs
    )


def _render_fixture(directory: Path, statement: str) -> tuple[Path, Path]:
    bullet = EvidenceBullet(
        fact_id="fact_service",
        raw_source_text=statement,
        source_kind="resume_extract",
        source_block_ids=["blk_service"],
        verified_statement=statement,
        user_confirmed=True,
    )
    bullet = bullet.model_copy(update={"content_sha256": content_sha256(bullet)})
    experience = WorkExperienceFact(
        exp_id="exp_service",
        company_name="Synthetic Systems",
        role_title="Engineer",
        start_date="2020",
        end_date="2022",
        bullets=[bullet],
        user_confirmed=True,
    )
    experience = experience.model_copy(update={"content_sha256": content_sha256(experience)})
    identity = Identity(user_confirmed=True)
    identity = identity.model_copy(update={"content_sha256": content_sha256(identity)})
    fact_base = UserFactBase(identity=identity, work_experiences=[experience])
    snapshot = FactBaseSnapshot(
        snapshot_id="snap_render_test",
        snapshot_sha256=sha256_json(fact_base),
        ledger_sha256="0" * 64,
        created_at="2026-01-01T00:00:00Z",
        fact_base=fact_base,
    )
    payload = TailoredResumePayload(
        snapshot_sha256=snapshot.snapshot_sha256,
        selected_skill_ids=[],
        experiences=[
            TailoredExperienceSection(
                exp_id=experience.exp_id,
                selected_bullets=[
                    RewrittenBullet(
                        bullet_id="bullet_rendered",
                        rewritten_text=statement,
                        source_fact_ids=[bullet.fact_id],
                        applied_style="balanced",
                    )
                ],
            )
        ],
    )
    snapshot_path = directory / "snapshot.json"
    payload_path = directory / "payload.json"
    snapshot_path.write_text(snapshot.model_dump_json(), encoding="utf-8")
    payload_path.write_text(payload.model_dump_json(), encoding="utf-8")
    return snapshot_path, payload_path


def test_markdown_and_html_fallback_are_literal_and_degraded(
    monkeypatch, tmp_path: Path
) -> None:
    statement = 'Built **reliable** service. ![x](http://example.test) <script>alert(1)</script>'
    snapshot_path, payload_path = _render_fixture(tmp_path, statement)
    output_dir = tmp_path / "candidate"
    monkeypatch.setattr("resume_tailor.rendering.shutil.which", lambda _: None)

    result, exit_code = render_resume(
        snapshot_path=snapshot_path,
        payload_path=payload_path,
        company_name=None,
        role_title=None,
        out_dir=output_dir,
        job_id=None,
        review_path=None,
        template_name="default",
        output_format="markdown",
        run_id=None,
    )
    markdown = (output_dir / "generated.md").read_text(encoding="utf-8")
    html_content = markdown_html(
        {
            "name": "",
            "contact": "",
            "summary": "",
            "skills": [],
            "experiences": [
                {
                    "role": "Engineer",
                    "company": "Synthetic Systems",
                    "dates": "2020 - 2022",
                    "bullets": [statement],
                }
            ],
            "projects": [],
            "educations": [],
        }
    )

    assert exit_code == 4
    assert result.status == "degraded"
    assert result.data["pdf_generated"] is False
    assert "generated.pdf" not in {Path(item.path).name for item in result.artifacts}
    assert r"\!\[x\]" in markdown
    assert "&lt;script&gt;" in markdown
    assert "<script>" not in html_content
    assert "<img" not in html_content
    assert "<a" not in html_content


def test_docx_render_degrades_without_libreoffice(monkeypatch, tmp_path: Path) -> None:
    snapshot_path, payload_path = _render_fixture(tmp_path, "Built a reliable service.")
    output_dir = tmp_path / "candidate"
    monkeypatch.setattr("resume_tailor.rendering.shutil.which", lambda _: None)

    result, exit_code = render_resume(
        snapshot_path=snapshot_path,
        payload_path=payload_path,
        company_name=None,
        role_title=None,
        out_dir=output_dir,
        job_id=None,
        review_path=None,
        template_name="default",
        output_format="docx",
        run_id=None,
    )

    assert exit_code == 4
    assert result.status == "degraded"
    assert (output_dir / "generated.docx").is_file()
    assert not (output_dir / "generated.pdf").exists()


def test_render_blocks_before_creating_output_when_fact_gate_fails(
    monkeypatch, tmp_path: Path
) -> None:
    snapshot_path, payload_path = _render_fixture(tmp_path, "Built a reliable service.")
    payload = TailoredResumePayload.model_validate_json(payload_path.read_text(encoding="utf-8"))
    payload.snapshot_sha256 = "f" * 64
    payload_path.write_text(payload.model_dump_json(), encoding="utf-8")
    output_dir = tmp_path / "blocked-candidate"
    monkeypatch.setattr("resume_tailor.rendering.shutil.which", lambda _: None)

    result, exit_code = render_resume(
        snapshot_path=snapshot_path,
        payload_path=payload_path,
        company_name=None,
        role_title=None,
        out_dir=output_dir,
        job_id=None,
        review_path=None,
        template_name="default",
        output_format="markdown",
        run_id=None,
    )

    assert exit_code == 3
    assert result.status == "blocked"
    assert not output_dir.exists()


def test_render_rechecks_job_requirements_before_creating_output(
    monkeypatch, tmp_path: Path
) -> None:
    snapshot_path, payload_path = _render_fixture(tmp_path, "Built a reliable service.")
    payload = TailoredResumePayload.model_validate_json(
        payload_path.read_text(encoding="utf-8")
    )
    payload.experiences[0].selected_bullets[0].rewritten_text = (
        "Used Kubernetes to operate the service."
    )
    payload_path.write_text(payload.model_dump_json(), encoding="utf-8")
    requirements = JobRequirements(
        jd_text_sha256="b" * 64,
        company_name="Synthetic Systems",
        role_title="Engineer",
        keywords=[JobKeyword(term="Kubernetes", category="skill", must_have=True)],
    )
    requirements_path = tmp_path / "requirements.json"
    requirements_path.write_text(requirements.model_dump_json(), encoding="utf-8")
    output_dir = tmp_path / "blocked-requirements-candidate"

    result, exit_code = render_resume(
        snapshot_path=snapshot_path,
        payload_path=payload_path,
        company_name=None,
        role_title=None,
        out_dir=output_dir,
        job_id=None,
        review_path=None,
        template_name="default",
        output_format="markdown",
        run_id=None,
        requirements_path=requirements_path,
    )

    issue_codes = {issue.code for issue in result.issues}
    assert exit_code == 3
    assert result.status == "blocked"
    assert "unsupported_term" in issue_codes
    assert "requirements_binding_mismatch" in issue_codes
    assert not output_dir.exists()


def test_render_health_failure_does_not_create_candidate_files(
    monkeypatch, tmp_path: Path
) -> None:
    snapshot_path, payload_path = _render_fixture(tmp_path, "Built a reliable service.")
    output_dir = tmp_path / "blocked-pdf-candidate"
    monkeypatch.setattr(
        "resume_tailor.rendering._convert_to_pdf",
        lambda source_name, source_bytes: b"synthetic-pdf",
    )
    monkeypatch.setattr(
        "resume_tailor.rendering.verify_pdf_health",
        lambda pdf_bytes, data, target_pages: [
            CheckResult(
                check="pdf_readable",
                passed=False,
                issue_code="pdf_unreadable",
                message="Generated PDF could not be inspected.",
            )
        ],
    )

    result, exit_code = render_resume(
        snapshot_path=snapshot_path,
        payload_path=payload_path,
        company_name=None,
        role_title=None,
        out_dir=output_dir,
        job_id=None,
        review_path=None,
        template_name="default",
        output_format="docx",
        run_id=None,
    )

    assert exit_code == 3
    assert result.status == "blocked"
    assert [issue.code for issue in result.issues] == ["pdf_unreadable"]
    assert not output_dir.exists()


@pytest.mark.parametrize("template_name", ["classic", "modern", "compact", "default"])
def test_builtin_templates_validate_and_render_successfully(template_name: str) -> None:
    from resume_tailor.rendering import _template_bytes, _validate_template

    template_bytes, name = _template_bytes(template_name)
    assert name == template_name
    assert len(template_bytes) > 0

    context = {
        "name": "Synthetic Candidate",
        "contact": "candidate@example.test | +1 (555) 010-0199",
        "summary": to_rich_text("Experienced synthetic engineer."),
        "skills": ["Python", "SQL", "Docker"],
        "experiences": [
            {
                "company": "Synthetic Corp",
                "role": "Lead Engineer",
                "dates": "2021 - Present",
                "bullets": [to_rich_text("Architected backend services.")],
            }
        ],
        "projects": [
            {
                "name": "Project Alpha",
                "role": "Creator",
                "bullets": [to_rich_text("Created scalable tools.")],
            }
        ],
        "educations": [{"label": "B.S. in Computer Science | Synthetic University | 2020"}],
    }

    _validate_template(template_bytes, context)
    rendered = render_docx_template(template_bytes, context)
    document = Document(io.BytesIO(rendered))
    full_text = " ".join(paragraph.text for paragraph in document.paragraphs)

    assert "Synthetic Candidate" in full_text
    assert "Synthetic Corp" in full_text
    assert "Architected backend services." in full_text


def test_find_soffice_detects_platform_default_paths(monkeypatch, tmp_path: Path) -> None:
    from resume_tailor.rendering import _find_soffice

    monkeypatch.setattr("resume_tailor.rendering.shutil.which", lambda _: None)
    monkeypatch.setattr("sys.platform", "win32")
    monkeypatch.setenv("ProgramFiles", str(tmp_path))

    # The default location expands to ProgramFiles / LibreOffice / program / soffice.exe
    mock_location = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    mock_location.parent.mkdir(parents=True, exist_ok=True)
    mock_location.touch()

    found = _find_soffice()
    assert found == str(mock_location)
